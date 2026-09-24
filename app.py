"""ScribeEase - extract text from handwritten/typed documents and export to .docx."""
import hmac
import io
import logging
import os
import re
import secrets
import time
from concurrent.futures import ThreadPoolExecutor

import pillow_heif
import pymupdf as fitz
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Inches
from dotenv import load_dotenv
from flask import (Flask, redirect, render_template, request, send_file,
                   session, url_for)
from google import genai
from google.genai import errors as genai_errors
from google.genai import types
from PIL import Image, ImageOps

pillow_heif.register_heif_opener()
load_dotenv()

# ---------- Configuration ----------

# Ordered list, best first. Free-tier daily quotas are counted PER MODEL, so when one model
# runs out (or is overloaded) the app moves on to the next. Flash-Lite has a much bigger free
# allowance but is less accurate, so it goes last.
DEFAULT_MODELS = ('gemini-3.5-flash,gemini-3.6-flash,gemini-3.7-flash,'
                  'gemini-3.8-flash,gemini-3.5-flash-lite')
MODELS = [m.strip() for m in os.environ.get('GEMINI_MODELS', DEFAULT_MODELS).split(',') if m.strip()]
QUOTA_COOLDOWN = 3600      # seconds to skip a model after it reports its daily limit reached
RETRY_DELAYS = (4, 12)                                       # seconds to wait before retries on a busy model
MAX_PAGES = int(os.environ.get('MAX_PAGES', '10'))          # total pages per request
MAX_UPLOAD_MB = int(os.environ.get('MAX_UPLOAD_MB', '25'))
ACCESS_CODE = os.environ.get('ACCESS_CODE', '')             # optional: gate the whole site
ALLOWED_EXTENSIONS = {'jpg', 'jpeg', 'png', 'webp', 'heic', 'heif',
                      'bmp', 'gif', 'tif', 'tiff', 'pdf'}

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = MAX_UPLOAD_MB * 1024 * 1024
app.secret_key = os.environ.get('SECRET_KEY') or secrets.token_hex(32)
logging.basicConfig(level=logging.INFO)

_client = None


def get_client():
    """Create the Gemini client on first use so a missing key gives a clear error, not a crash loop."""
    global _client
    if _client is None:
        key = os.environ.get('GEMINI_API_KEY')
        if not key:
            raise UserError('The server is missing its GEMINI_API_KEY setting.')
        _client = genai.Client(api_key=key)
    return _client


class UserError(Exception):
    """An error whose message is safe and useful to show to the person using the site."""
    status = 400


class ServiceBusy(UserError):
    status = 503


class QuotaExhausted(ServiceBusy):
    pass


EXTRACTION_PROMPT = (
    'IMPORTANT FORMATTING RULE: Ignore the visual line breaks in the image caused by the writer running out of space at the edge of the page. Reflow the text into natural paragraphs. '
    'Only start a new line in your output when there is a genuine blank line/gap in the image, a new numbered item, a new question, or a clear topic change — never simply because a physical line of handwriting or print ended. '
    'For example, if the image shows "Angelina aquired land from her friend on" on one line and "a temporary basis to grow maize." on the next line with no gap between them, output this as one single continuous sentence: "Angelina aquired land from her friend on a temporary basis to grow maize." — do not keep them as two separate lines. '
    'Extract all text from this image exactly as written, whether handwritten or typed/printed, following the reflow rule above. '
    'For mathematical and chemical notation, mark superscript text using ^{...} and subscript text using _{...} — for example, x^{2} for "x squared" and H_{2}O for water. '
    'Preserve arrows, Greek letters, and simple symbols exactly as they appear using standard Unicode characters (e.g. →, Δ, ∫, ±, θ, ÷). '
    'For square roots, write as sqrt(...) — for example sqrt(x^{2} + 1). '
    'For fractions, write as (numerator)/(denominator) — for example (x+1)/(x-1). '
    'For binomial coefficients, write as C(n,r) or n choose r. '
    'For chemical equations, use → for reactions and ⇌ for reversible/equilibrium reactions. Write state symbols as (s), (l), (g), or (aq) immediately after each substance, not as subscript — for example NaCl(aq) not NaCl_{aq}. Preserve reaction conditions written above or below the arrow by placing them in brackets after the arrow, like →[heat] or →[catalyst, 200°C]. '
    'If the image contains a table, format it as a Markdown table using | to separate columns and a header separator row of dashes (e.g. | Header 1 | Header 2 |\\n|---|---|\\n| data | data |). '
    'When two pieces of text appear on the same line but positioned at opposite ends (like "Time allowed" on the left and "Maximum Marks" on the right), separate them with a single Tab character so their positions are preserved. '
    'Return only the extracted text, no commentary.'
)


# ---------- DOCX helpers ----------

def add_runs_with_formatting(paragraph, line):
    line = line.replace('<br>', '\n').replace('<br/>', '\n').replace('<BR>', '\n')

    if '\t' in line:
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)

    segments = line.split('\n')
    pattern = r'(\^\{[^}]*\}|_\{[^}]*\})'
    for seg_index, segment in enumerate(segments):
        if seg_index > 0:
            paragraph.add_run().add_break()

        tab_parts = segment.split('\t')
        for t_index, tab_part in enumerate(tab_parts):
            if t_index > 0:
                paragraph.add_run('\t')
            parts = re.split(pattern, tab_part)
            for part in parts:
                if not part:
                    continue
                if part.startswith('^{') and part.endswith('}'):
                    run = paragraph.add_run(part[2:-1])
                    run.font.superscript = True
                elif part.startswith('_{') and part.endswith('}'):
                    run = paragraph.add_run(part[2:-1])
                    run.font.subscript = True
                else:
                    paragraph.add_run(part)


def add_formatted_paragraph(doc, line):
    paragraph = doc.add_paragraph()
    add_runs_with_formatting(paragraph, line)


def is_table_separator(line):
    stripped = line.strip()
    if not (stripped.startswith('|') and stripped.endswith('|')):
        return False
    inner = stripped.strip('|')
    parts = inner.split('|')
    return all(re.match(r'^\s*:?-+:?\s*$', p) for p in parts)


def parse_table_row(line):
    stripped = line.strip()
    if stripped.startswith('|'):
        stripped = stripped[1:]
    if stripped.endswith('|'):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split('|')]


def add_table_to_doc(doc, rows):
    if not rows:
        return
    num_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j in range(num_cols):
            cell_text = row[j] if j < len(row) else ''
            cell_paragraph = table.cell(i, j).paragraphs[0]
            add_runs_with_formatting(cell_paragraph, cell_text)


# ---------- Reading uploads (everything stays in memory; nothing is written to disk) ----------

def image_to_jpeg_bytes(img):
    img = ImageOps.exif_transpose(img)      # phone photos: respect the rotation flag
    img = img.convert('RGB')
    img.thumbnail((2600, 2600))             # keep payloads small; handwriting stays legible
    buf = io.BytesIO()
    img.save(buf, 'JPEG', quality=90)
    return buf.getvalue()


def pdf_to_jpeg_pages(data, limit):
    try:
        doc = fitz.open(stream=data, filetype='pdf')
    except Exception:
        raise UserError('One of the PDFs could not be opened.')
    try:
        if doc.needs_pass:
            raise UserError('Password-protected PDFs are not supported.')
        if doc.page_count > limit:
            raise UserError(f'Too many pages. The limit is {MAX_PAGES} pages per upload.')
        return [page.get_pixmap(dpi=200).tobytes('jpeg') for page in doc]
    finally:
        doc.close()


def collect_pages(files):
    """Turn every uploaded file into a flat list of JPEG page images."""
    pages = []
    for f in files:
        name = f.filename or ''
        ext = name.rsplit('.', 1)[-1].lower() if '.' in name else ''
        if ext not in ALLOWED_EXTENSIONS:
            raise UserError(f'"{name}" is not a supported file type. Use JPG, PNG, HEIC, WebP or PDF.')
        data = f.read()
        if not data:
            raise UserError(f'"{name}" is empty.')
        if ext == 'pdf':
            pages.extend(pdf_to_jpeg_pages(data, MAX_PAGES - len(pages)))
        else:
            try:
                pages.append(image_to_jpeg_bytes(Image.open(io.BytesIO(data))))
            except Exception:
                raise UserError(f'"{name}" could not be read as an image.')
        if len(pages) > MAX_PAGES:
            raise UserError(f'Too many pages. The limit is {MAX_PAGES} pages per upload.')
    return pages


# ---------- Gemini ----------

def extract_text_from_page(jpeg_bytes, model):
    response = get_client().models.generate_content(
        model=model,
        contents=[types.Part.from_bytes(data=jpeg_bytes, mime_type='image/jpeg'),
                  EXTRACTION_PROMPT],
    )
    return (response.text or '').strip() or '[No text could be extracted from this page]'


def is_transient(exc):
    """Overloaded / rate-limited / temporary server trouble: worth waiting and retrying."""
    return isinstance(exc, genai_errors.APIError) and getattr(exc, 'code', None) in (429, 500, 502, 503, 504)


def is_daily_quota(exc):
    """The per-day request cap (free tier: ~20 per Flash model). Retrying can't help until the day resets."""
    return (isinstance(exc, genai_errors.APIError) and getattr(exc, 'code', None) == 429
            and 'PerDay' in str(exc))


_exhausted_until = {}      # model -> time.time() until which we skip it (daily quota used up)


def extract_with_retry(item):
    """Returns (text, model_used), or None if every model and retry failed.

    Each round tries every usable model once, then waits a little longer before the next round.
    Overload (503) and quota (429) are per-model, so a struggling model is quickly swapped for the next."""
    number, jpeg_bytes = item
    now = time.time()
    usable = [m for m in MODELS if _exhausted_until.get(m, 0) <= now]
    if not usable:
        raise QuotaExhausted('Today\'s usage limit for the AI service has been reached. '
                             'Please try again later (the free limit resets around 10-11 AM Uganda time).')
    out_of_quota = set()
    for delay in (0,) + RETRY_DELAYS:
        if delay:
            time.sleep(delay)
        for model in list(usable):
            try:
                return extract_text_from_page(jpeg_bytes, model), model
            except UserError:
                raise
            except Exception as e:
                app.logger.warning('Page %s, model %s failed: %s', number, model, e)
                if is_daily_quota(e):
                    out_of_quota.add(model)
                    _exhausted_until[model] = time.time() + QUOTA_COOLDOWN
                    usable.remove(model)   # don't waste more requests on a model that is out for today
                elif not is_transient(e):
                    usable.remove(model)   # e.g. unknown model: retrying it won't help
        if not usable:
            break
    if out_of_quota and not usable:
        raise QuotaExhausted('Today\'s usage limit for the AI service has been reached. '
                             'Please try again later (the free limit resets around 10-11 AM Uganda time).')
    app.logger.error('Giving up on page %s', number)
    return None


# ---------- Access gate (optional) ----------

@app.before_request
def require_access_code():
    if not ACCESS_CODE or request.endpoint in ('login', 'static'):
        return None
    if not session.get('authorised'):
        return redirect(url_for('login'))
    return None


@app.route('/login', methods=['GET', 'POST'])
def login():
    if not ACCESS_CODE:
        return redirect(url_for('index'))
    error = None
    if request.method == 'POST':
        supplied = request.form.get('code', '')
        if hmac.compare_digest(supplied.encode(), ACCESS_CODE.encode()):
            session['authorised'] = True
            return redirect(url_for('index'))
        time.sleep(1)  # slow down guessing
        error = 'Incorrect access code.'
    return render_template('login.html', error=error), (401 if error else 200)


# ---------- Routes ----------

@app.route('/')
def index():
    return render_template('index.html', max_pages=MAX_PAGES, max_mb=MAX_UPLOAD_MB)


@app.route('/healthz')
def healthz():
    return 'ok'


@app.route('/upload', methods=['POST'])
def upload():
    files = [f for f in request.files.getlist('documents') if f.filename]
    try:
        if not files:
            raise UserError('Please choose at least one file.')
        pages = collect_pages(files)
        with ThreadPoolExecutor(max_workers=3) as pool:
            results = list(pool.map(extract_with_retry, enumerate(pages, start=1)))
        if all(r is None for r in results):
            raise ServiceBusy('The AI service is busy or unavailable right now. Nothing was lost: '
                              'please try again in a few minutes.')
        page_texts = [r[0] if r is not None else f'[Page {i} could not be read. Please upload it again.]'
                      for i, r in enumerate(results, start=1)]
        used = {r[1] for r in results if r is not None}
        notice = None
        if any('lite' in m for m in used) and 'lite' not in MODELS[0]:   # only warn if Lite was a fallback
            notice = ('Some pages were read by a lighter backup model because the main models were busy '
                      'or out of daily quota. Please check them extra carefully against the original.')
    except UserError as e:
        return render_template('index.html', error=str(e),
                               max_pages=MAX_PAGES, max_mb=MAX_UPLOAD_MB), e.status

    if len(page_texts) == 1:
        combined = page_texts[0]
    else:
        combined = ''
        for i, text in enumerate(page_texts, start=1):
            if i > 1:
                combined += f'\n\n--- Page {i} ---\n\n'
            combined += text
    return render_template('result.html', text=combined, notice=notice)


PAGE_MARKER = re.compile(r'^--- Page \d+ ---$')


@app.route('/download', methods=['POST'])
def download():
    # Browsers send textarea line breaks as \r\n; a stray \r becomes an extra line break in Word.
    text = request.form.get('edited_text', '').replace('\r\n', '\n').replace('\r', '\n')

    doc = Document()
    lines = text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        if PAGE_MARKER.match(line.strip()):
            doc.add_page_break()
            i += 1
            while i < len(lines) and not lines[i].strip():   # skip blank lines after the marker
                i += 1
        elif line.strip().startswith('|') and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            table_lines = [line]
            i += 2  # skip the header row and the separator row
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            add_table_to_doc(doc, [parse_table_row(l) for l in table_lines])
            doc.add_paragraph('')
        else:
            add_formatted_paragraph(doc, line)
            i += 1

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return send_file(
        stream,
        as_attachment=True,
        download_name='ScribeEase_output.docx',
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )


@app.errorhandler(413)
def too_large(_e):
    return render_template('index.html', error=f'Upload too large. The limit is {MAX_UPLOAD_MB} MB in total.',
                           max_pages=MAX_PAGES, max_mb=MAX_UPLOAD_MB), 413


if __name__ == '__main__':
    app.run(debug=os.environ.get('FLASK_DEBUG') == '1')
